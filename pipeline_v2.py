"""v2.0 pipeline — R1 新结构 + R2b 独立提取 spots
用法: python pipeline_v2.py <docx_path>
"""
import asyncio
import json
import os
import sys
import time
from pathlib import Path

from dotenv import load_dotenv
from openai import AsyncOpenAI
from openpyxl import Workbook

import urllib.request
import urllib.parse

from prompts_v2 import make_system, R1_USER, R2A_USER, R2B_USER, build_r2c_user
from filter import filter_skeleton_for_r2
from run_layer1 import load_docx

load_dotenv(Path(__file__).parent / ".env")

R1_SCHEMA = {
    "type": "object",
    "properties": {
        "reference": {
            "type": "object",
            "properties": {
                "title": {"type": "string"},
                "authors": {"type": "string"},
                "journal": {"type": "string"},
                "year": {"type": "string"},
                "volume": {"type": "string"},
                "pages": {"type": "string"},
                "doi": {"type": "string"}
            },
            "required": ["title", "authors", "journal", "year", "volume", "pages", "doi"],
            "additionalProperties": False
        },
        "samples": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "id": {"type": "string"},
                    "origin": {"type": "string"},
                    "rock_type": {"type": "string"},
                    "formation": {"type": "string"},
                    "group": {"type": "string"},
                    "member": {"type": "string"},
                    "observations": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "id": {"type": "string"},
                                "sample_material": {
                                    "anyOf": [
                                        {"type": "null"},
                                        {"type": "string"}
                                    ]
                                },
                                "measurements": {
                                    "type": "array",
                                    "items": {
                                        "type": "object",
                                        "properties": {
                                            "property": {"type": "string"},
                                            "age": {
                                                "anyOf": [
                                                    {"type": "null"},
                                                    {"type": "object", "properties": {
                                                        "number": {"anyOf": [{"type": "null"}, {"type": "number"}]},
                                                        "error": {"type": "object", "properties": {
                                                            "value": {"anyOf": [{"type": "null"}, {"type": "number"}]},
                                                            "scale": {"type": "string"}
                                                        }, "required": ["value", "scale"], "additionalProperties": False}
                                                    }, "required": ["number", "error"], "additionalProperties": False}
                                                ]
                                            },
                                            "ratio": {
                                                "anyOf": [
                                                    {"type": "null"},
                                                    {"type": "object", "properties": {
                                                        "number": {"anyOf": [{"type": "null"}, {"type": "number"}]},
                                                        "error": {"type": "object", "properties": {
                                                            "value": {"anyOf": [{"type": "null"}, {"type": "number"}]},
                                                            "scale": {"type": "string"}
                                                        }, "required": ["value", "scale"], "additionalProperties": False}
                                                    }, "required": ["number", "error"], "additionalProperties": False}
                                                ]
                                            }
                                        },
                                        "required": ["property", "age", "ratio"],
                                        "additionalProperties": False
                                    }
                                }
                            },
                            "required": ["id", "sample_material", "measurements"],
                            "additionalProperties": False
                        }
                    },
                    "measure_events": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "method": {"type": "string"},
                                "technique": {"type": "string"},
                                "value": {"type": "object", "properties": {
                                    "number": {"anyOf": [{"type": "null"}, {"type": "number"}]},
                                    "error": {"type": "object", "properties": {
                                        "value": {"anyOf": [{"type": "null"}, {"type": "number"}]},
                                        "scale": {"type": "string"}
                                    }, "required": ["value", "scale"], "additionalProperties": False}
                                }, "required": ["number", "error"], "additionalProperties": False},
                                "distribution_state": {"anyOf": [{"type": "null"}, {"type": "string"}]},
                                "statistical_model": {"anyOf": [{"type": "null"}, {"type": "string"}]},
                                "annotation": {"anyOf": [{"type": "null"}, {"type": "string"}]}
                            },
                            "required": ["method", "technique", "value", "distribution_state", "statistical_model", "annotation"],
                            "additionalProperties": False
                        }
                    }
                },
                "required": ["id", "origin", "rock_type", "formation", "group", "member", "observations", "measure_events"],
                "additionalProperties": False
            }
        }
    },
    "required": ["reference", "samples"],
    "additionalProperties": False
}

API_KEY = os.environ["DEEPSEEK_API_KEY"]
BASE_URL = os.environ["DEEPSEEK_BASE_URL"]
MODEL = os.environ["DEEPSEEK_MODEL"]
OPENALEX_EMAIL = os.environ.get("OPENALEX_EMAIL", "structured-extract@example.com")

# R2b spot 允许的字段
_SPOT_KEYS = {"id", "u_ppm", "th_ppm", "th_u_ratio", "concordance",
              "r207_206", "r207_235", "r206_238", "age_206_238", "age_207_235"}
# R2a 只修改这些槽位
_R2A_KEYS = {"rock_type", "formation", "group", "member",
             "province", "city", "basin", "locality"}
# R2c 受控词表字段
_EVENT_BASE_KEYS = {"value", "method", "annotation"}
_EVENT_VOCAB_KEYS = {"technique", "mineral", "isotope_system", "physical_method"}


def _log_cache(usage, label: str):
    hit = getattr(usage, 'prompt_cache_hit_tokens', 0) or 0
    miss = getattr(usage, 'prompt_cache_miss_tokens', 0) or 0
    total = hit + miss
    pct = f"{hit/total*100:.0f}%" if total else "N/A"
    print(f"  [{label}] cache: hit={hit} miss={miss} ({pct})")


def enrich_reference(ref: dict) -> dict:
    title = (ref.get("title") or "").strip()
    if not title or len(title) < 10:
        return ref
    filled = {k: v for k, v in ref.items() if v}
    if all(filled.get(k) for k in ("doi", "journal", "year", "volume", "pages")):
        return ref
    q = urllib.parse.quote(title[:150])
    url = f"https://api.openalex.org/works?search={q}&per_page=1&mailto={OPENALEX_EMAIL}"
    try:
        req = urllib.request.Request(url, headers={
            "User-Agent": f"structured-extract/1.0 (mailto:{OPENALEX_EMAIL})"})
        resp = urllib.request.urlopen(req, timeout=15)
        data = json.loads(resp.read())
        results = data.get("results", [])
        if not results:
            print("  OpenAlex: 无匹配")
            return ref
        w = results[0]
        loc = w.get("primary_location", {}) or {}
        source = loc.get("source") or {}
        if not filled.get("doi"):
            raw_doi = w.get("doi") or ""
            filled["doi"] = raw_doi.replace("https://doi.org/", "")
        if not filled.get("journal"):
            filled["journal"] = source.get("display_name", "")
        if not filled.get("year"):
            y = w.get("publication_year")
            filled["year"] = str(y) if y else ""
        if not filled.get("volume"):
            bib = w.get("biblio") or {}
            filled["volume"] = bib.get("volume") or ""
        if not filled.get("pages"):
            bib = w.get("biblio") or {}
            first = bib.get("first_page") or ""
            last = bib.get("last_page") or ""
            filled["pages"] = f"{first}-{last}" if first and last else (first or "")
        added = [k for k in ("doi", "journal", "year", "volume", "pages") if filled.get(k) and not ref.get(k)]
        if added:
            print(f"  OpenAlex 补充: {', '.join(added)}")
        return filled
    except Exception as e:
        print(f"  OpenAlex 查询失败: {str(e)[:80]}")
        return ref


def _clean_r2a_output(r1_sample: dict, r2a_sample: dict) -> dict:
    out = dict(r1_sample)
    for k in _R2A_KEYS:
        out[k] = r2a_sample.get(k, "")
    if "location" in r2a_sample:
        out["location"] = r2a_sample["location"]
    return out


def _clean_spot(spot: dict) -> dict:
    return {k: spot.get(k) for k in _SPOT_KEYS if k in spot}


def _clean_r2c_output(r1_sample: dict, r2c_sample: dict) -> list:
    r1_events = r1_sample.get("measure_events", [])
    r2c_events = r2c_sample.get("measure_events", [])
    cleaned = []
    for i, r1e in enumerate(r1_events):
        out = {k: r1e[k] for k in _EVENT_BASE_KEYS if k in r1e}
        r2ce = r2c_events[i] if i < len(r2c_events) else {}
        r1_val = r1e.get("value") or {}
        r1_err = r1_val.get("error") or {}
        r2c_val = r2ce.get("value") or {}
        r2c_err = r2c_val.get("error") or {}
        out["value"] = {
            "number": r1_val.get("number"),
            "error": {
                "value": r1_err.get("value"),
                "scale": r2c_err.get("scale", r1_err.get("scale", "未声明")),
            },
        }
        out["statistical_model"] = r2ce.get("statistical_model", r1e.get("statistical_model", "null"))
        out["distribution_state"] = r2ce.get("distribution_state", r1e.get("distribution_state", "null"))
        for k in _EVENT_VOCAB_KEYS:
            out[k] = r2ce.get(k, "null")
        cleaned.append(out)
    return cleaned


# ═══════════════════════════════════════════════════
# Round 1（v2.0 新结构）
# ═══════════════════════════════════════════════════

async def run_r1(client: AsyncOpenAI, paper_text: str) -> dict:
    t0 = time.time()
    resp = await client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": make_system(paper_text)},
            {"role": "user", "content": R1_USER},
        ],
        temperature=0,
        reasoning_effort="max",
        extra_body={"thinking": {"type": "enabled"}},
        response_format={"type": "json_object"},
    )
    elapsed = time.time() - t0
    usage = resp.usage
    print(f"[R1] {elapsed:.1f}s | in={usage.prompt_tokens} out={usage.completion_tokens}")
    _log_cache(usage, "R1")
    return json.loads(resp.choices[0].message.content)


# ═══════════════════════════════════════════════════
# Round 2a（沿用 v1）
# ═══════════════════════════════════════════════════

async def run_r2a_one(client: AsyncOpenAI, paper_text: str, sample: dict) -> dict:
    skel = {"samples": [sample]}
    skel_json = json.dumps(skel, ensure_ascii=False)
    try:
        resp = await client.chat.completions.create(
            model=MODEL,
            messages=[
                {"role": "system", "content": make_system(paper_text)},
                {"role": "user", "content": f"=== Entity Tree JSON ===\n{skel_json}\n\n{R2A_USER}"},
            ],
            temperature=0,
            reasoning_effort="max",
            extra_body={"thinking": {"type": "enabled"}},
            response_format={"type": "json_object"},
        )
    except Exception as e:
        print(f"  [ERROR] R2a:{sample['id']} {str(e)[:100]}")
        return sample
    usage = resp.usage
    _log_cache(usage, f"R2a:{sample['id']}")
    result = json.loads(resp.choices[0].message.content)
    if "samples" not in result:
        print(f"  [WARN] R2a:{sample['id']} 返回无 samples")
        return sample
    return _clean_r2a_output(sample, result["samples"][0])


# ═══════════════════════════════════════════════════
# Round 2b（v2.0: 独立提取 spots，不依赖 R1 observations）
# ═══════════════════════════════════════════════════

R2B_V2_USER = """任务：从文本/表格中为每个样品提取原始观测点（spots），填入对应 samples 的 spots 子节点。

═══ 待填槽位 ═══

- spots[].id — 观测点标识符，从表格中逐字抄录
- u_ppm, th_ppm, th_u_ratio, concordance — 直接填数值或 null
- r207_206, r207_235, r206_238 — 比值字段，结构为 {value, err_1s}
- age_206_238, age_207_235 — 年龄字段，结构为 {value, err_1s}

═══ 规则 ═══

1. 从表格中逐行读取每个样品的所有观测点数据
2. spots[].id 必须与表格中的标识符完全一致
3. 数值直接从表格读取，不计算、不推测
4. 表格中缺失的数值填 null
5. 仅填充 spots，其他槽位原样保留
6. 输出完整 JSON"""


async def run_r2b_v2(client: AsyncOpenAI, paper_text: str, sample_ids: list[str]) -> dict:
    """v2 R2b: 传入 sample id 列表，独立从文本提取 spots。"""
    skeleton = {"samples": [{"id": sid, "spots": []} for sid in sample_ids]}
    skel_json = json.dumps(skeleton, ensure_ascii=False)
    try:
        resp = await client.chat.completions.create(
            model=MODEL,
            messages=[
                {"role": "system", "content": make_system(paper_text)},
                {"role": "user", "content": f"=== Entity Tree JSON ===\n{skel_json}\n\n{R2B_V2_USER}"},
            ],
            temperature=0,
            reasoning_effort="max",
            extra_body={"thinking": {"type": "enabled"}},
            response_format={"type": "json_object"},
        )
    except Exception as e:
        print(f"  [ERROR] R2b {str(e)[:100]}")
        return {}
    usage = resp.usage
    _log_cache(usage, "R2b")
    result = json.loads(resp.choices[0].message.content)
    if "samples" not in result:
        print(f"  [WARN] R2b 返回无 samples")
        return {}
    return {s["id"]: s for s in result["samples"]}


# ═══════════════════════════════════════════════════
# Round 2c（沿用 v1）
# ═══════════════════════════════════════════════════

async def run_r2c_one(client: AsyncOpenAI, paper_text: str, sample: dict) -> list:
    skel = {"samples": [sample]}
    skel_json = json.dumps(skel, ensure_ascii=False)
    resp = await client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": make_system(paper_text)},
            {"role": "user", "content": build_r2c_user(skel_json)},
        ],
        temperature=0,
        reasoning_effort="max",
        extra_body={"thinking": {"type": "enabled"}},
        response_format={"type": "json_object"},
    )
    usage = resp.usage
    _log_cache(usage, f"R2c:{sample['id']}")
    result = json.loads(resp.choices[0].message.content)
    if "samples" not in result:
        print(f"  [WARN] R2c:{sample['id']} 返回无 samples")
        return _clean_r2c_output(sample, {"measure_events": [{}] * len(sample.get("measure_events", []))})
    return _clean_r2c_output(sample, result["samples"][0])


# ═══════════════════════════════════════════════════
# Excel 导出（沿用 v1）
# ═══════════════════════════════════════════════════

_SAMPLE_COLS = ["sample_id", "origin", "rock_type", "formation", "group", "member",
                "province", "city", "basin", "locality",
                "lat", "lng", "location_source"]
_EVENT_COLS = ["technique", "value_number", "value_error_value", "value_error_scale",
               "annotation", "isotope_system", "mineral", "physical_method",
               "statistical_model", "distribution_state"]
_SPOT_COLS = ["u_ppm", "th_ppm", "th_u_ratio", "concordance",
              "r207_206_value", "r207_206_err",
              "r207_235_value", "r207_235_err",
              "r206_238_value", "r206_238_err",
              "age_206_238_value", "age_206_238_err",
              "age_207_235_value", "age_207_235_err"]


def _spot_row(sample_id: str, spot: dict) -> list:
    row = [sample_id]
    for k in _SPOT_COLS:
        if k.endswith(("_value", "_err")):
            rk, sub = k.rsplit("_", 1)
            sub = "err_1s" if sub == "err" else sub
            obj = spot.get(rk)
            row.append(obj.get(sub) if isinstance(obj, dict) else None)
        else:
            row.append(spot.get(k))
    return ["" if v in (None, "null") else v for v in row]


def _event_row(e: dict) -> list:
    val = e.get("value") or {}
    err = val.get("error") or {}
    row = []
    for k in _EVENT_COLS:
        if k == "value_number":
            row.append(val.get("number"))
        elif k == "value_error_value":
            row.append(err.get("value"))
        elif k == "value_error_scale":
            row.append(err.get("scale"))
        elif k == "statistical_model":
            row.append(e.get("statistical_model"))
        elif k == "distribution_state":
            row.append(e.get("distribution_state"))
        else:
            row.append(e.get(k))
    return ["" if v in (None, "null") else v for v in row]


def _export_flat_xlsx(merged: dict, out_path: Path):
    wb = Workbook()
    ws_ref = wb.active
    ws_ref.title = "reference"
    ref = merged.get("reference", {})
    for k in ["title", "authors", "journal", "year", "volume", "pages", "doi"]:
        ws_ref.append([k, ref.get(k) or ""])
    ws_e = wb.create_sheet("events")
    ws_e.append(_SAMPLE_COLS + _EVENT_COLS)
    for s in merged.get("samples", []):
        loc = s.get("location") or {}
        base = [s["id"], s.get("origin", ""), s.get("rock_type", ""),
                s.get("formation", ""), s.get("group", ""), s.get("member", ""),
                s.get("province", ""), s.get("city", ""), s.get("basin", ""),
                s.get("locality", ""),
                loc.get("lat"), loc.get("lng"), loc.get("source", "")]
        for e in s.get("measure_events", []):
            ws_e.append(base + _event_row(e))
    ws_s = wb.create_sheet("spots")
    ws_s.append(["sample_id"] + _SPOT_COLS)
    for s in merged.get("samples", []):
        for sp in s.get("spots", []):
            ws_s.append(_spot_row(s["id"], sp))
    wb.save(out_path)


def _export_tree_xlsx(merged: dict, out_path: Path):
    wb = Workbook()
    ws_ref = wb.active
    ws_ref.title = "reference"
    ref = merged.get("reference", {})
    for k in ["title", "authors", "journal", "year", "volume", "pages", "doi"]:
        ws_ref.append([k, ref.get(k) or ""])
    ws_s = wb.create_sheet("samples")
    sample_cols = ["sample_id", "origin", "rock_type", "formation", "group", "member",
                   "province", "city", "basin", "locality", "lat", "lng", "location_source"]
    ws_s.append(sample_cols)
    for s in merged.get("samples", []):
        loc = s.get("location") or {}
        row = [s["id"], s.get("origin", ""), s.get("rock_type", ""),
               s.get("formation", ""), s.get("group", ""), s.get("member", ""),
               s.get("province", ""), s.get("city", ""), s.get("basin", ""),
               s.get("locality", ""),
               loc.get("lat"), loc.get("lng"), loc.get("source", "")]
        ws_s.append(["" if v in (None, "null") else v for v in row])
    ws_sp = wb.create_sheet("spots")
    spot_cols = ["sample_id", "spot_id", "u_ppm", "th_ppm", "th_u_ratio", "concordance",
                 "r207_206_value", "r207_206_err",
                 "r207_235_value", "r207_235_err",
                 "r206_238_value", "r206_238_err",
                 "age_206_238_value", "age_206_238_err",
                 "age_207_235_value", "age_207_235_err"]
    ws_sp.append(spot_cols)
    for s in merged.get("samples", []):
        for sp in s.get("spots", []):
            row = _spot_row(s["id"], sp)
            row.insert(1, sp.get("id", ""))
            ws_sp.append(["" if v in (None, "null") else v for v in row])
    ws_e = wb.create_sheet("events")
    event_cols = ["sample_id", "technique", "value_number", "value_error_value",
                  "value_error_scale", "annotation", "isotope_system", "mineral",
                  "physical_method", "statistical_model", "distribution_state"]
    ws_e.append(event_cols)
    for s in merged.get("samples", []):
        for e in s.get("measure_events", []):
            row = [s["id"]] + _event_row(e)
            ws_e.append(["" if v in (None, "null") else v for v in row])
    wb.save(out_path)


# ═══════════════════════════════════════════════════
# v2.0 Pipeline
# ═══════════════════════════════════════════════════

async def run_pipeline(docx_path: str):
    paper_text = load_docx(docx_path)
    client = AsyncOpenAI(api_key=API_KEY, base_url=BASE_URL)
    out_dir = Path(__file__).parent / "output"
    out_dir.mkdir(exist_ok=True)
    stem = Path(docx_path).parent.name or Path(docx_path).stem

    # ── R1（v2.0 新结构）──
    print(f"文件: {Path(docx_path).name} ({len(paper_text)} 字符)")
    t_total = time.time()
    skeleton = await run_r1(client, paper_text)
    n_samples = len(skeleton.get("samples", []))
    n_events = sum(len(s.get("measure_events", [])) for s in skeleton.get("samples", []))
    n_obs = sum(len(s.get("observations", [])) for s in skeleton.get("samples", []))
    print(f"  R1 结果: samples={n_samples} obs={n_obs} events={n_events}")
    (out_dir / f"{stem}_r1.json").write_text(
        json.dumps(skeleton, ensure_ascii=False, indent=2))
    print(f"  已保存: {stem}_r1.json")

    # ── OpenAlex 补全 reference ──
    skeleton["reference"] = enrich_reference(skeleton.get("reference", {}))

    # ── Filter（只过滤 cited samples）──
    n_before = len(skeleton.get("samples", []))
    skeleton = filter_skeleton_for_r2(skeleton)
    n_after = len(skeleton.get("samples", []))
    print(f"  过滤: {n_after} self samples ({n_before - n_after} cited removed)")

    # ── R2a + R2b + R2c 并发 ──
    self_samples = skeleton.get("samples", [])
    sample_ids = [s["id"] for s in self_samples]
    n = len(self_samples)
    print(f"\n  R2 启动: R2a×{n} + R2b×1 + R2c×{n}")
    t_r2 = time.time()

    # 缓存预热：第一个 R2a 顺序跑
    first_sample = self_samples[0]
    print(f"  缓存预热: R2a:{first_sample['id']}...", end=" ", flush=True)
    first_r2a = await run_r2a_one(client, paper_text, first_sample)
    print("done")

    # 剩余并发
    tasks_r2a = [run_r2a_one(client, paper_text, s) for s in self_samples[1:]]
    tasks_r2b = [run_r2b_v2(client, paper_text, sample_ids)]
    tasks_r2c = [run_r2c_one(client, paper_text, s) for s in self_samples]

    results_2a_rest, results_2b_map, results_2c_raw = await asyncio.gather(
        asyncio.gather(*tasks_r2a),
        asyncio.gather(*tasks_r2b),
        asyncio.gather(*tasks_r2c),
    )
    results_2a = [first_r2a] + list(results_2a_rest)
    r2b_map = results_2b_map[0]
    r2c_map = {self_samples[i]["id"]: results_2c_raw[i] for i in range(n)}

    print(f"  R2 完成: {time.time() - t_r2:.1f}s")

    # ── 合并：R1 events + R2a attrs + R2b spots + R2c vocab ──
    merged = {"reference": skeleton.get("reference", {}), "samples": []}
    for r2a_sample in results_2a:
        sid = r2a_sample["id"]
        sample = dict(r2a_sample)
        # R2b spots
        if sid in r2b_map:
            r2b_spots = r2b_map[sid].get("spots", [])
            sample["spots"] = [_clean_spot(sp) for sp in r2b_spots]
        else:
            sample["spots"] = []
        # R2c events（用 R1 的 events 做 base）
        if sid in r2c_map:
            sample["measure_events"] = r2c_map[sid]
        # 删除 R1 的 observations（脚手架，不进最终输出）
        sample.pop("observations", None)
        merged["samples"].append(sample)

    # ── 保存 ──
    out_path = out_dir / f"{stem}_final.json"
    out_path.write_text(json.dumps(merged, ensure_ascii=False, indent=2))
    xlsx_path = out_dir / f"{stem}_flat.xlsx"
    _export_flat_xlsx(merged, xlsx_path)
    tree_path = out_dir / f"{stem}_tree.xlsx"
    _export_tree_xlsx(merged, tree_path)
    print(f"\n  总耗时: {time.time() - t_total:.1f}s")
    print(f"  保存: {out_path}")
    print(f"  保存: {xlsx_path}")
    print(f"  保存: {tree_path}")


def main():
    if len(sys.argv) < 2:
        print("用法: python pipeline_v2.py <docx_path>")
        sys.exit(1)
    asyncio.run(run_pipeline(sys.argv[1]))


if __name__ == "__main__":
    main()
