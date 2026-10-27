"""多篇论文合并导出 — 从多个 _final.json 生成一个 total_tree.xlsx"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from pipeline import (
    _is_noise_event, _is_empty_event, _dedup_events,
    _event_row, _spot_row,
)

# ═══════════════════════════════════════════════════
# 列定义（左侧加 paper_id）
# ═══════════════════════════════════════════════════

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(
    bottom=Side(style="thin", color="D9D9D9"),
)

REF_COLS = ["paper_id", "title", "authors", "journal", "year", "volume", "pages", "doi"]

SAMPLE_COLS = ["paper_id", "sample_id", "origin", "rock_type", "formation", "group",
               "member", "province", "city", "basin", "locality", "lat", "lng"]

OBS_COLS = ["paper_id", "sample_id", "spot_id", "u_ppm", "th_ppm", "th_u_ratio",
            "concordance", "r207_206_value", "r207_206_err", "r207_235_value",
            "r207_235_err", "r206_238_value", "r206_238_err", "age_206_238_value",
            "age_206_238_err", "age_207_235_value", "age_207_235_err"]

EVENT_COLS = ["paper_id", "sample_id", "technique", "value_number", "value_error_value",
              "value_error_scale", "annotation", "isotope_system", "mineral",
              "physical_method", "statistical_model", "distribution_state"]

_SPOT_NUMERIC_KEYS = [
    "u_ppm", "th_ppm", "th_u_ratio", "concordance",
    "r207_206", "r207_235", "r206_238", "age_206_238", "age_207_235",
]


def _safe(v):
    return "" if v in (None, "null") else v


def _write_header(ws, cols):
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _spot_row_padded(spot: dict) -> list:
    """从 spot dict 提取 OBS_COLS[3:] 对应的数值列。"""
    row = []
    for k in _SPOT_NUMERIC_KEYS:
        if k in ("r207_206", "r207_235", "r206_238", "age_206_238", "age_207_235"):
            obj = spot.get(k)
            if isinstance(obj, dict):
                row.append(_safe(obj.get("value")))
                row.append(_safe(obj.get("err_1s")))
            else:
                row.extend(["", ""])
        else:
            v = spot.get(k)
            if isinstance(v, dict):
                v = v.get("value")
            row.append(_safe(v))
    return row


def _event_row_padded(e: dict) -> list:
    """从 event dict 提取 EVENT_COLS[2:] 对应的字段。"""
    val = e.get("value") or {}
    err = val.get("error") or {}
    return [
        _safe(e.get("technique")),
        _safe(val.get("number")),
        _safe(err.get("value")),
        _safe(err.get("scale")),
        _safe(e.get("annotation")),
        _safe(e.get("isotope_system")),
        _safe(e.get("mineral")),
        _safe(e.get("physical_method")),
        _safe(val.get("statistical_model")),
        _safe(val.get("distribution_state")),
    ]


def _is_empty_spot(spot: dict) -> bool:
    return all(
        spot.get(k) in (None, "null", "")
        for k in _SPOT_NUMERIC_KEYS
    )


def export_total(json_dir: str | Path, out_path: str | Path = None):
    json_dir = Path(json_dir)
    if out_path is None:
        out_path = json_dir / "total_tree.xlsx"
    else:
        out_path = Path(out_path)

    files = sorted(json_dir.glob("*_final.json"))
    if not files:
        print("无 *_final.json 文件")
        return

    wb = Workbook()
    ws_ref = wb.active
    ws_ref.title = "reference"
    ws_sam = wb.create_sheet("samples")
    ws_obs = wb.create_sheet("observations")
    ws_evt = wb.create_sheet("measure_events")

    _write_header(ws_ref, REF_COLS)
    _write_header(ws_sam, SAMPLE_COLS)
    _write_header(ws_obs, OBS_COLS)
    _write_header(ws_evt, EVENT_COLS)

    n_papers = 0
    for fp in files:
        data = json.loads(fp.read_text())
        paper_id = data.get("reference", {}).get("doi") or fp.stem.replace("_final", "")
        ref = data.get("reference", {})

        # reference
        ws_ref.append([_safe(paper_id)] + [_safe(ref.get(k)) for k in REF_COLS[1:]])

        for s in data.get("samples", []):
            sid = s["id"]
            loc = s.get("location") or {}
            ws_sam.append([
                _safe(paper_id), _safe(sid), _safe(s.get("origin")),
                _safe(s.get("rock_type")), _safe(s.get("formation")),
                _safe(s.get("group")), _safe(s.get("member")),
                _safe(s.get("province")), _safe(s.get("city")),
                _safe(s.get("basin")), _safe(s.get("locality")),
                loc.get("lat"), loc.get("lng"),
            ])

            # spots
            for sp in s.get("spots", []):
                if _is_empty_spot(sp):
                    continue
                ws_obs.append([_safe(paper_id), _safe(sid), _safe(sp.get("id"))]
                              + _spot_row_padded(sp))

            # events（过滤 + 去重）
            events = [e for e in s.get("measure_events", [])
                      if not _is_noise_event(e) and not _is_empty_event(e)]
            for e in _dedup_events(events):
                ws_evt.append([_safe(paper_id), _safe(sid)] + _event_row_padded(e))

        n_papers += 1

    wb.save(out_path)
    print(f"已导出: {out_path}  ({n_papers} 篇论文)")


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("用法: python export_total.py <json_dir> [out.xlsx]")
    else:
        export_total(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
