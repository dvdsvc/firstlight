"""多份 tree.xlsx 合并 — 从多个独立 tree.xlsx 汇总成一个 total_tree.xlsx
用法: python merge_tree_xlsx.py <dir1> [dir2 ...] [-o out.xlsx]
"""
import argparse
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

EXPECTED_SHEETS = ["reference", "samples", "observations", "measure_events"]
REFERENCE_COLUMNS = ["paper_id", "title", "authors", "journal", "year", "volume", "pages", "doi"]
COORDINATE_COLUMNS = [
    "lat_normalized",
    "lng_normalized",
    "lat_geocoded",
    "lng_geocoded",
    "coordinate_status",
    "coordinate_source",
]


def _paper_id_from_path(fp: Path) -> str:
    # The filename is the stable source identity. Do not use DOI here:
    # different source files can have the same, missing, or incorrect DOI.
    return fp.stem[:-5] if fp.stem.endswith("_tree") else fp.stem


def _write_header(ws, cols):
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _read_reference(ws) -> dict:
    """Read the per-paper reference sheet written by pipeline._export_tree_xlsx.

    Those source sheets are key/value rows without a header, e.g.
    ("title", "...") and ("authors", "..."). Also accept an already-wide
    reference sheet so the merge remains usable with mixed source versions.
    """
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}

    first = [str(value).strip() if value is not None else "" for value in rows[0]]
    if first and first[0] in {"paper_id", "title"} and len(first) > 2:
        header = first
        values = list(rows[1]) if len(rows) > 1 else []
        return {
            key: values[index] if index < len(values) else ""
            for index, key in enumerate(header)
            if key
        }

    return {
        str(row[0]).strip(): row[1] if len(row) > 1 else ""
        for row in rows
        if row and row[0] is not None and str(row[0]).strip()
    }


def _read_tab_rows(ws) -> tuple[list[str], list[dict]]:
    """Read a tab with its first row as the column header."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    header = [str(value).strip() if value is not None else "" for value in rows[0]]
    records = []
    for row in rows[1:]:
        values = list(row)
        while len(values) < len(header):
            values.append("")
        records.append({key: values[index] for index, key in enumerate(header) if key})
    return header, records


def _normalize_lat_lng(record: dict) -> tuple[float, float] | None:
    """从分开的 lat/lng 读取原生坐标，不改变 record 中的原值。"""
    lat_value = record.get("lat")
    lng_value = record.get("lng")
    if lat_value in (None, "") or lng_value in (None, ""):
        return None
    try:
        lat = float(str(lat_value).strip())
        lng = float(str(lng_value).strip())
    except (TypeError, ValueError):
        return None
    if abs(lat) > 90 and abs(lng) <= 90:
        lat, lng = lng, lat
    if abs(lat) > 90 or abs(lng) > 180 or (lat == 0 and lng == 0):
        return None
    return round(lat, 6), round(lng, 6)


def merge_tree_xlsx(dirs: list[Path], out_path: Path):
    files = []
    for d in dirs:
        d = Path(d)
        if d.is_file() and d.suffix == ".xlsx":
            files.append(d)
        elif d.is_dir():
            files.extend(sorted(d.glob("*_tree.xlsx")))
    if not files:
        print("无 *_tree.xlsx 文件")
        return

    wb_out = Workbook()
    ws_ref = wb_out.active
    ws_ref.title = "reference"
    _write_header(ws_ref, REFERENCE_COLUMNS)

    ws_map = {"reference": ws_ref}
    header_map = {"reference": REFERENCE_COLUMNS}
    n = 0

    for fp in files:
        try:
            wb = load_workbook(fp, read_only=True, data_only=True)
        except Exception as e:
            print(f"  跳过 {fp.name}: {e}")
            continue

        sheets_found = [s for s in EXPECTED_SHEETS if s in wb.sheetnames]
        if not sheets_found:
            print(f"  跳过 {fp.name}: 无标准 sheet")
            continue

        paper_id = _paper_id_from_path(fp)

        # reference 的源格式是 key/value，而不是带表头的数据表。
        if "reference" in sheets_found:
            ref = _read_reference(wb["reference"])
            ws_ref.append([paper_id] + [ref.get(column, "") for column in REFERENCE_COLUMNS[1:]])

        for sname in (name for name in sheets_found if name != "reference"):
            orig_header, records = _read_tab_rows(wb[sname])
            if not orig_header:
                continue

            # 以第一份输入的列顺序作为输出契约；后续文件按列名对齐，
            # 不再按位置硬拼，避免不同版本 tree 表发生错列。
            if sname not in ws_map:
                source_columns = [
                    key for key in orig_header
                    if key not in {"paper_id", "GPS"} and key not in COORDINATE_COLUMNS
                ]
                output_header = ["paper_id"] + source_columns
                if sname == "samples":
                    # lat/lng 是原生输入；其余列分别保存解析结果和地理编码结果。
                    output_header.extend(column for column in COORDINATE_COLUMNS if column not in output_header)
                ws = wb_out.create_sheet(sname)
                ws_map[sname] = ws
                header_map[sname] = output_header
                _write_header(ws, output_header)

            ws = ws_map[sname]
            output_header = header_map[sname]
            for record in records:
                row = [paper_id]
                for column in output_header[1:]:
                    value = record.get(column, "")
                    if sname == "samples" and column in COORDINATE_COLUMNS:
                        normalized = _normalize_lat_lng(record)
                        if column == "lat_normalized" and value in (None, "") and normalized:
                            value = normalized[0]
                        elif column == "lng_normalized" and value in (None, "") and normalized:
                            value = normalized[1]
                        elif column == "coordinate_status" and value in (None, ""):
                            value = "valid_native" if normalized else "missing_or_invalid_native"
                        elif column == "coordinate_source" and value in (None, "") and normalized:
                            value = "native"
                    row.append(value)
                ws.append(row)

        wb.close()
        n += 1

    wb_out.save(out_path)
    print(f"已合并 {n} 篇 → {out_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("dirs", nargs="+", help="目录或 xlsx 文件")
    parser.add_argument("-o", "--out", default=None, help="输出路径")
    args = parser.parse_args()

    out = Path(args.out) if args.out else Path(args.dirs[0]) / "total_tree.xlsx" if len(args.dirs) == 1 else Path("total_tree.xlsx")
    merge_tree_xlsx([Path(d) for d in args.dirs], out)
