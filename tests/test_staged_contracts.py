"""Synthetic integration and adverse-response characterization of original code.

These cases establish local contracts and limitations, not real-paper accuracy.
"""
from copy import deepcopy
from contextlib import redirect_stdout, redirect_stderr
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from examples.staged_demo import event, fixture, sample, read_table, run_demo
from scripts.offline_runtime import load_staged, no_network
from scripts.export_linked import export

P = load_staged()
V2 = load_staged('pipeline_v2')


class StagedContracts(unittest.TestCase):
    def test_full_offline_orchestration_and_paper_links(self):
        self.assertTrue(run_demo()['same_sample_id_across_two_papers_kept_separate'])

    def test_r2a_field_boundary_and_missing_attribute_behavior(self):
        s = sample()
        s['rock_type'] = 'old'
        result = P._clean_r2a_output(s, {'id': 'OTHER', 'spots': [], 'arbitrary': 3})
        self.assertEqual(result['id'], s['id'])
        self.assertEqual(result['spots'], s['spots'])
        self.assertEqual(result['rock_type'], '')  # Missing approved attrs become empty.
        self.assertNotIn('arbitrary', result)

    def test_r2c_protects_r1_number_error_and_annotation(self):
        s = sample()
        e = event(999)
        e['value']['error']['value'] = 999
        e['annotation'] = 'do not replace'
        result = P._clean_r2c_output(s, {'measure_events': [e]})[0]
        self.assertEqual(result['value']['number'], 120)
        self.assertEqual(result['value']['error']['value'], 1)
        self.assertEqual(result['annotation'], s['measure_events'][0]['annotation'])

    def test_r2c_is_positional_and_does_not_validate_vocabulary(self):
        s = sample()
        s['measure_events'].append(event(240))
        a, b = event(240), event(120)
        a['technique'], b['technique'] = 'UNKNOWN-TERM', 'LA-ICPMS'
        result = P._clean_r2c_output(s, {'measure_events': [a,b]})
        self.assertEqual(result[0]['value']['number'], 120)
        self.assertEqual(result[0]['technique'], 'UNKNOWN-TERM')

    def test_r2b_has_field_allowlist_not_strict_id_membership(self):
        s = sample()
        result = P._clean_r2b_output(s, {'spots': [{'id': 'DEMO-NEW', 'u_ppm': 8, 'unknown': 9}]})
        self.assertEqual(result, [{'id': 'DEMO-NEW', 'u_ppm': 8}])
        self.assertEqual(P._clean_r2b_output(s, {'spots': []}), [])

    def test_r1_union_keeps_base_sample_set_and_numeric_dedup(self):
        base = fixture()
        other = fixture()
        other['samples'][0]['measure_events'][0]['method'] = 'DIFFERENT'
        other['samples'].append({**sample(), 'id': 'DEMO-NEW'})
        result = P.merge_r1_events(base, [other])
        self.assertNotIn('DEMO-NEW', {s['id'] for s in result['samples']})
        self.assertEqual(len(result['samples'][0]['measure_events']), 1)

    def test_filter_copies_and_repairs_numeric_suffix(self):
        d = fixture()
        before = deepcopy(d)
        result = P.filter_skeleton_for_r2(d)
        self.assertEqual(d, before)
        self.assertEqual(len(result['samples']), 1)
        self.assertEqual(result['samples'][0]['spots'][0]['id'], 'DEMO-SAME-1')
        self.assertNotIn('method', result['samples'][0]['measure_events'][0])

    def test_export_projection_changes_rows_without_mutating_json(self):
        d = fixture()
        d['samples'] = d['samples'][:1]
        e = event()
        e['value']['statistical_model'] = 'Weighted average age'
        d['samples'][0]['measure_events'] = [deepcopy(e) for _ in range(5)] + [event(240)]
        before = deepcopy(d)
        with tempfile.TemporaryDirectory() as tmp, no_network():
            out = Path(tmp) / 'DEMO_tree.xlsx'
            P._export_tree_xlsx(d, out)
            wb = load_workbook(out, data_only=True)
            self.assertEqual(len(read_table(wb, 'measure_events')), 1)
            wb.close()
        self.assertEqual(d, before)
        self.assertEqual(len(d['samples'][0]['measure_events']), 6)

    def test_empty_export_may_produce_no_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'empty.xlsx'
            P._export_tree_xlsx({'reference': {}, 'samples': []}, path)
            self.assertFalse(path.exists())

    def test_duplicate_paper_ids_rejected_before_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for n in ('A', 'B'):
                (root / f'{n}_final.json').write_text(json.dumps(fixture()))
            with self.assertRaisesRegex(ValueError, 'duplicate paper_id'):
                export(root, root / 'out.xlsx')
            self.assertFalse((root / 'out.xlsx').exists())

    def test_export_wrapper_refuses_overwrite(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'A_final.json').write_text(json.dumps(fixture()))
            target = root / 'out.xlsx'
            target.write_bytes(b'existing artifact')
            with self.assertRaises(FileExistsError):
                export(root, target)
            self.assertEqual(target.read_bytes(), b'existing artifact')

    def test_export_wrapper_rejects_v2_event_schema(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            d = fixture()
            d['samples'][0]['measure_events'][0]['statistical_model'] = 'Weighted average age'
            (root / 'A_final.json').write_text(json.dumps(d))
            with self.assertRaisesRegex(ValueError, 'matching v2 exporter'):
                export(root, root / 'out.xlsx')
            self.assertFalse((root / 'out.xlsx').exists())

    def test_live_wrapper_blocks_existing_outputs_for_both_variants(self):
        from scripts import run_staged
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'scripts').mkdir()
            (root / 'output').mkdir()
            (root / 'DEMO').mkdir()
            document = root / 'DEMO/full.docx'
            document.write_bytes(b'fixture input; must never be opened')
            target = root / 'output/DEMO_final.json'
            target.write_bytes(b'existing output')
            for variant in ('main','v2'):
                with self.subTest(variant=variant), patch.object(run_staged,'__file__',str(root/'scripts/run_staged.py')), \
                     patch('sys.argv',['run_staged',str(document),'--variant',variant]), \
                     patch.object(run_staged.importlib,'import_module',side_effect=AssertionError('Must stop before research imports')), \
                     redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                    with self.assertRaises(SystemExit):
                        run_staged.main()
            self.assertEqual(target.read_bytes(), b'existing output')

    def test_current_tree_merger_reads_key_value_reference_and_aligns_columns(self):
        merger = load_staged('merge_tree_xlsx')
        with tempfile.TemporaryDirectory() as tmp, no_network():
            root = Path(tmp)
            d = fixture()
            d['samples'] = d['samples'][:1]
            d['samples'][0]['spots'][0]['u_ppm'] = 12
            for name in ('A','B'):
                P._export_tree_xlsx(d, root / f'{name}_tree.xlsx')
            b = load_workbook(root / 'B_tree.xlsx')
            ws = b['samples']
            # Deliberately reorder two source columns in the second workbook.
            a_vals, b_vals = [c.value for c in ws['A']], [c.value for c in ws['B']]
            for i, (av,bv) in enumerate(zip(a_vals,b_vals),1):
                ws.cell(i,1,bv)
                ws.cell(i,2,av)
            b.save(root / 'B_tree.xlsx')
            b.close()
            with redirect_stdout(io.StringIO()):
                merger.merge_tree_xlsx([root], root / 'total.xlsx')
            wb = load_workbook(root / 'total.xlsx', data_only=True)
            self.assertEqual({r['paper_id'] for r in read_table(wb,'reference')}, {'A','B'})
            self.assertTrue(all(r['title'] == 'Fictional DEMO paper A' for r in read_table(wb,'reference')))
            self.assertTrue(all(r['sample_id'] == 'DEMO-SAME' and r['origin'] == 'self' for r in read_table(wb,'samples')))
            wb.close()

    def test_v2_keeps_statistics_at_event_top_level(self):
        s = sample()
        e = event(999)
        e['statistical_model'] = 'Weighted average age'
        result = V2._clean_r2c_output(s, {'measure_events': [e]})[0]
        self.assertEqual(result['value']['number'], 120)
        self.assertEqual(result['statistical_model'], 'Weighted average age')
        self.assertNotIn('statistical_model', result['value'])


if __name__ == '__main__':
    unittest.main()
