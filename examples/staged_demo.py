"""Execute the original staged orchestrator using handwritten fake responses.

The DOCX, values and responses are fictional. No LLM inference is performed.
"""
import asyncio
from copy import deepcopy
from contextlib import redirect_stdout
import io
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
from unittest.mock import patch
from docx import Document
from openpyxl import load_workbook
from scripts.offline_runtime import load_staged, no_network
from scripts.export_linked import export


def event(number=120.0):
    return {'method': 'U-Pb', 'technique': 'LA-ICPMS',
            'value': {'number': number, 'error': {'value': 1.0, 'scale': '2σ'}},
            'annotation': 'Fictional DEMO sample age.'}


def sample():
    return {'id': 'DEMO-SAME', 'origin': 'self', 'rock_type': '', 'formation': '',
            'group': '', 'member': '', 'spots': [{'id': '1'}],
            'measure_events': [event()]}


def fixture():
    s = sample()
    return {'reference': {'title': 'Fictional DEMO paper A', 'doi': 'DEMO-PAPER-A'},
            'samples': [s, {**deepcopy(s), 'id': 'DEMO-CITED', 'origin': 'cited'}]}


class FixtureClient:
    def __init__(self, **kwargs):
        self.chat = SimpleNamespace(completions=self)
        self.calls = 0

    async def create(self, **kwargs):
        self.calls += 1
        user = kwargs['messages'][-1]['content']
        if 'Entity Tree JSON' not in user:
            answer = fixture()
        elif 'ontology annotation' in user:
            e = event(9999)  # The original merge must retain the R1 numeric value.
            e.update({'technique': 'LA-ICPMS', 'mineral': 'Zircon',
                      'isotope_system': 'U-Th-Pb dating', 'physical_method': 'Ablation / Spot'})
            e['value']['error']['value'] = 999
            e['value'].update({'statistical_model': 'Weighted average age', 'distribution_state': 'null'})
            answer = {'samples': [{**sample(), 'measure_events': [e]}]}
        elif 'spots 子节点' in user:
            answer = {'samples': [{**sample(), 'spots': [{'id': 'DEMO-SAME-1', 'u_ppm': 12,
                       'th_ppm': 6, 'age_206_238': {'value': 120, 'err_1s': 1},
                       'unregistered_column': 'must not survive'}]}]}
        else:
            answer = {'samples': [{**sample(), 'rock_type': 'Fictional tuff',
                       'location': {'lat': 30, 'lng': 110, 'source': 'paper'}}]}
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(answer)))],
                               usage=SimpleNamespace(prompt_tokens=0, completion_tokens=0))


def read_table(book, name):
    rows = list(book[name].values)
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def run_demo():
    module = load_staged()
    with tempfile.TemporaryDirectory(prefix='structured-extract-demo-') as tmp, no_network():
        root = Path(tmp)
        inp = root / 'DEMO-A' / 'full.docx'
        inp.parent.mkdir()
        doc = Document()
        doc.add_paragraph('Fictional DEMO document; sample DEMO-SAME has an age of 120 ± 1 Ma.')
        table = doc.add_table(rows=1, cols=3)
        for c, value in zip(table.rows[0].cells, ['DEMO-SAME-1', '120', '1']):
            c.text = value
        doc.save(inp)
        client = FixtureClient()
        with redirect_stdout(io.StringIO()), patch.object(module, '__file__', str(root / 'pipeline.py')), \
             patch.object(module, 'AsyncOpenAI', return_value=client), \
             patch.object(module, 'enrich_reference', side_effect=lambda r: r):
            asyncio.run(module.run_pipeline(str(inp)))
        final = root / 'output/DEMO-A_final.json'
        data = json.loads(final.read_text())
        assert len(data['samples']) == 1
        s = data['samples'][0]
        assert s['id'] == 'DEMO-SAME' and s['rock_type'] == 'Fictional tuff'
        assert s['measure_events'][0]['value']['number'] == 120
        assert s['measure_events'][0]['value']['error']['value'] == 1
        assert s['spots'][0]['id'] == 'DEMO-SAME-1' and 'unregistered_column' not in s['spots'][0]
        second = deepcopy(data)
        second['reference'] = {'title': 'Fictional DEMO paper B', 'doi': 'DEMO-PAPER-B'}
        second['samples'][0]['measure_events'][0]['value']['number'] = 240
        (final.parent / 'DEMO-B_final.json').write_text(json.dumps(second))
        linked = root / 'linked.xlsx'
        with redirect_stdout(io.StringIO()):
            export(final.parent, linked)
        wb = load_workbook(linked, data_only=True)
        events = read_table(wb, 'measure_events')
        observed = {(r['paper_id'], r['sample_id']): r['value_number'] for r in events}
        assert observed == {('DEMO-PAPER-A', 'DEMO-SAME'): 120, ('DEMO-PAPER-B', 'DEMO-SAME'): 240}
        counts = {name: wb[name].max_row - 1 for name in wb.sheetnames}
        wb.close()
        return {'data_kind': 'fictional fixture responses, not model predictions',
                'network_requests': 0, 'model_calls': 0, 'fixture_calls': client.calls,
                'original_orchestrator_executed': True, 'retained_self_samples': 1,
                'r1_numeric_value_preserved_against_r2c_overwrite': True,
                'unknown_spot_field_removed': True, 'linked_rows': counts,
                'same_sample_id_across_two_papers_kept_separate': True,
                'temporary_outputs_removed_after_demo': True}


if __name__ == '__main__':
    print(json.dumps(run_demo(), indent=2))
