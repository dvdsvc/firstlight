"""Round 1 骨架抽取 — 用法: python run_layer1.py <docx_path>"""
import json
import os
import sys
import time
from pathlib import Path

from dotenv import load_dotenv
from docx import Document
from openai import OpenAI

from prompts import make_system, R1_USER

load_dotenv(Path(__file__).parent / ".env")

API_KEY = os.environ["DEEPSEEK_API_KEY"]
BASE_URL = os.environ["DEEPSEEK_BASE_URL"]
MODEL = os.environ["DEEPSEEK_MODEL"]


def load_docx(path: str) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for i, table in enumerate(doc.tables):
        parts.append(f"\n[表格 {i}]")
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def build_r1_schema() -> dict:
    return {
        "type": "object",
        "properties": {
            "reference": {
                "type": "object",
                "properties": {
                    "title": {"type": "string"},
                    "authors": {"type": "string"},
                    "journal": {"type": "string"},
                    "year": {"type": "string"},
                    "volume": {"type": "string"},
                    "pages": {"type": "string"},
                    "doi": {"type": "string"},
                },
                "required": ["title", "authors", "journal", "year", "volume", "pages", "doi"],
            },
            "samples": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "id": {"type": "string"},
                        "origin": {"type": "string"},
                        "rock_type": {"type": "string"},
                        "formation": {"type": "string"},
                        "group": {"type": "string"},
                        "member": {"type": "string"},
                        "location": {
                            "type": "object",
                            "properties": {
                                "lat": {"type": ["number", "null"]},
                                "lng": {"type": ["number", "null"]},
                                "source": {"type": "string"},
                            },
                            "required": ["lat", "lng", "source"],
                        },
                        "spots": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "properties": {
                                    "id": {"type": "string"},
                                    "u_ppm": {"type": ["number", "null"]},
                                    "th_ppm": {"type": ["number", "null"]},
                                },
                                "required": ["id"],
                            },
                        },
                        "measure_events": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "properties": {
                                    "method": {"type": "string"},
                                    "technique": {"type": "string"},
                                    "value": {
                                        "type": "object",
                                        "properties": {
                                            "number": {"type": ["number", "null"]},
                                            "error": {
                                                "type": "object",
                                                "properties": {
                                                    "value": {"type": ["number", "null"]},
                                                    "scale": {"type": "string"},
                                                },
                                                "required": ["value", "scale"],
                                            },
                                        },
                                        "required": ["number", "error"],
                                    },
                                    "annotation": {"type": "string"},
                                },
                                "required": [
                                    "method", "technique", "value", "annotation",
                                ],
                            },
                        },
                    },
                    "required": [
                        "id", "origin", "rock_type", "formation",
                        "group", "member", "location", "spots", "measure_events",
                    ],
                },
            }
        },
        "required": ["samples"],
    }


def extract_skeleton(text: str) -> dict:
    client = OpenAI(api_key=API_KEY, base_url=BASE_URL)
    t0 = time.time()
    resp = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": make_system(text)},
            {"role": "user", "content": R1_USER},
        ],
        temperature=0,
        reasoning_effort="max",
        extra_body={"thinking": {"type": "enabled"}},
        response_format={"type": "json_object"},
    )
    elapsed = time.time() - t0
    content = resp.choices[0].message.content
    usage = resp.usage

    data = json.loads(content)
    print(f"  耗时 {elapsed:.1f}s | tokens: in={usage.prompt_tokens} out={usage.completion_tokens}")
    print(f"  缓存命中: {getattr(usage, 'prompt_cache_hit_tokens', 'N/A')} / 未命中: {getattr(usage, 'prompt_cache_miss_tokens', 'N/A')}")
    return data


def main():
    if len(sys.argv) < 2:
        print("用法: python run_layer1.py <docx_path>")
        sys.exit(1)

    docx_path = sys.argv[1]
    print(f"文件: {Path(docx_path).name}")

    text = load_docx(docx_path)
    print(f"  文本长度: {len(text)} 字符")

    result = extract_skeleton(text)

    n_samples = len(result.get("samples", []))
    n_events = sum(len(s.get("measure_events", [])) for s in result.get("samples", []))
    n_spots = sum(len(s.get("spots", [])) for s in result.get("samples", []))
    print(f"  抽取结果: samples={n_samples} events={n_events} spots={n_spots}")

    out_dir = Path(__file__).parent / "output"
    out_dir.mkdir(exist_ok=True)
    stem = Path(docx_path).stem
    out_path = out_dir / f"{stem}_layer1.json"
    out_path.write_text(json.dumps(result, ensure_ascii=False, indent=2))
    print(f"  保存: {out_path}")


if __name__ == "__main__":
    main()
