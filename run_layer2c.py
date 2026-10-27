"""Round 2c 事件分类 — 用法: python run_layer2c.py <docx_path> <layer1_json_path>"""
import json
import os
import sys
import time
from pathlib import Path

from dotenv import load_dotenv
from openai import OpenAI

from prompts import make_system, build_r2c_user
from filter import filter_skeleton_for_r2
from vocab import (
    ISOTOPE_SYSTEM_OPTIONS, DATING_TECHNIQUE_OPTIONS,
    DATING_MINERAL_OPTIONS,
    PHYSICAL_METHOD_OPTIONS, STATISTICAL_MODEL_OPTIONS,
    DISTRIBUTION_STATE_OPTIONS, ERROR_SCALE_OPTIONS,
)

load_dotenv(Path(__file__).parent / ".env")

API_KEY = os.environ["DEEPSEEK_API_KEY"]
BASE_URL = os.environ["DEEPSEEK_BASE_URL"]
MODEL = os.environ["DEEPSEEK_MODEL"]


def _enum_field(enum_values: list[str], description: str) -> dict:
    return {"type": "string", "enum": enum_values, "description": description}


def build_r2c_schema() -> dict:
    event_schema = {
        "type": "object",
        "properties": {
            "value": {
                "type": "object",
                "properties": {
                    "number": {"type": ["number", "null"]},
                    "error": {
                        "type": "object",
                        "properties": {
                            "value": {"type": ["number", "null"]},
                            "scale": _enum_field(ERROR_SCALE_OPTIONS, "误差标度"),
                        },
                        "required": ["value", "scale"],
                    },
                    "statistical_model": _enum_field(
                        STATISTICAL_MODEL_OPTIONS + ["null"], "统计模型"),
                    "distribution_state": _enum_field(
                        DISTRIBUTION_STATE_OPTIONS + ["null"], "分布状态"),
                },
                "required": ["number", "error", "statistical_model", "distribution_state"],
            },
            "annotation": {"type": "string"},
            "isotope_system": _enum_field(ISOTOPE_SYSTEM_OPTIONS, "同位素体系"),
            "technique": _enum_field(DATING_TECHNIQUE_OPTIONS, "测年技术"),
            "mineral": _enum_field(DATING_MINERAL_OPTIONS, "测定矿物"),
            "physical_method": _enum_field(PHYSICAL_METHOD_OPTIONS + ["null"], "物理方法"),
        },
        "required": [
            "value", "annotation",
            "isotope_system", "technique", "mineral",
            "physical_method",
        ],
    }

    return {
        "type": "object",
        "properties": {
            "samples": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "id": {"type": "string"},
                        "origin": {"type": "string"},
                        "rock_type": {"type": "string"},
                        "formation": {"type": "string"},
                        "group": {"type": "string"},
                        "member": {"type": "string"},
                        "province": {"type": "string"},
                        "city": {"type": "string"},
                        "basin": {"type": "string"},
                        "locality": {"type": "string"},
                        "location": {
                            "type": "object",
                            "properties": {
                                "lat": {"type": ["number", "null"]},
                                "lng": {"type": ["number", "null"]},
                                "source": {"type": "string"},
                            },
                            "required": ["lat", "lng", "source"],
                        },
                        "spots": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "properties": {"id": {"type": "string"}},
                                "required": ["id"],
                            },
                        },
                        "measure_events": {
                            "type": "array",
                            "items": event_schema,
                        },
                    },
                    "required": [
                        "id", "origin", "rock_type", "formation",
                        "group", "member", "location", "spots", "measure_events",
                    ],
                },
            }
        },
        "required": ["samples"],
    }


def load_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for i, table in enumerate(doc.tables):
        parts.append(f"\n[表格 {i}]")
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def classify_events(paper_text: str, skeleton: dict) -> dict:
    client = OpenAI(api_key=API_KEY, base_url=BASE_URL)
    skel_json = json.dumps(skeleton, ensure_ascii=False)
    t0 = time.time()
    resp = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": make_system(paper_text)},
            {"role": "user", "content": build_r2c_user(skel_json)},
        ],
        temperature=0,
        reasoning_effort="max",
        extra_body={"thinking": {"type": "enabled"}},
        response_format={"type": "json_object"},
    )
    elapsed = time.time() - t0
    content = resp.choices[0].message.content
    usage = resp.usage
    data = json.loads(content)
    print(f"  耗时 {elapsed:.1f}s | tokens: in={usage.prompt_tokens} out={usage.completion_tokens}")
    print(f"  缓存命中: {getattr(usage, 'prompt_cache_hit_tokens', 'N/A')} / 未命中: {getattr(usage, 'prompt_cache_miss_tokens', 'N/A')}")
    return data


def main():
    if len(sys.argv) < 3:
        print("用法: python run_layer2c.py <docx_path> <layer1_json_path>")
        sys.exit(1)

    docx_path, skel_path = sys.argv[1], sys.argv[2]
    paper_text = load_docx(docx_path)
    skeleton = json.loads(Path(skel_path).read_text())

    n_before = len(skeleton.get("samples", []))
    skeleton = filter_skeleton_for_r2(skeleton)
    n_after = len(skeleton.get("samples", []))
    print(f"文件: {Path(docx_path).name}")
    print(f"  骨架: {n_after} self samples (过滤 {n_before - n_after} cited)")

    result = classify_events(paper_text, skeleton)

    n_events = sum(len(s.get("measure_events", [])) for s in result.get("samples", []))
    print(f"  分类结果: {n_events} events")

    out_dir = Path(__file__).parent / "output"
    out_dir.mkdir(exist_ok=True)
    stem = Path(docx_path).stem
    out_path = out_dir / f"{stem}_layer2c.json"
    out_path.write_text(json.dumps(result, ensure_ascii=False, indent=2))
    print(f"  保存: {out_path}")


if __name__ == "__main__":
    main()
